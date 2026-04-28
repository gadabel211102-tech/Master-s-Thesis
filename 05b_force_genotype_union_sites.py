#!/usr/bin/env python3
"""
Script 05b: Force genotyping across all cohort-observed targeted loci.

Purpose
-------
Build a cohort-wide catalogue of all loci observed in the filtered targeted VCF
set, then force-genotype those loci across every sample BAM. This complements
script 05 by distinguishing true wild-type / reference genotypes from loci that
were simply not emitted by the original variant caller.

Methodological note
-------------------
This stage does not replace the main high-confidence variant call set. Instead,
it provides a genotype-centric view across the union of targeted loci that were
observed at least once anywhere in the cohort. That means it can recover 0/0,
0/1, 1/1 and no-call states for cohort-observed loci, while keeping the main
PASS-filtered call set unchanged for discovery analyses.
"""

from __future__ import annotations

import argparse
import math
import re
import shlex
import subprocess
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pipeline_utils import get_paths, normalise_chromosome_label, standardize_cohort_labels, standardize_tissue_labels

PATHS = get_paths()
DEFAULT_OUTPUT_DIR = PATHS.get("forced_genotypes_dir", PATHS["results_dir"] / "05b_forced_genotypes")

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=12)
WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="top")
STATUS_FILLS = {
    "WT": PatternFill(fill_type="solid", fgColor="D9EAF7"),
    "Het_ALT": PatternFill(fill_type="solid", fgColor="FFF2CC"),
    "Hom_ALT": PatternFill(fill_type="solid", fgColor="FCE4D6"),
    "Below_DP_threshold": PatternFill(fill_type="solid", fgColor="FDE9D9"),
    "No_call": PatternFill(fill_type="solid", fgColor="E7E6E6"),
    "Other_ALT_model": PatternFill(fill_type="solid", fgColor="E4DFEC"),
}


def log(message: str) -> None:
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{timestamp}] {message}")


def ensure_directory(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def ensure_tool(tool_name: str) -> None:
    if shutil_which(tool_name) is None:
        raise RuntimeError(f"Required tool was not found on PATH: {tool_name}")


def shutil_which(tool_name: str) -> str | None:
    from shutil import which

    return which(tool_name)


def run_command(command: list[str], description: str) -> subprocess.CompletedProcess[str]:
    try:
        return subprocess.run(command, check=True, capture_output=True, text=True)
    except subprocess.CalledProcessError as exc:
        stderr = (exc.stderr or "").strip()
        stdout = (exc.stdout or "").strip()
        detail = stderr or stdout or "No stderr/stdout captured."
        raise RuntimeError(f"{description} failed: {' '.join(command)}\n{detail}") from exc


def chrom_sort_key(chrom: str) -> tuple[int, str]:
    norm = normalise_chromosome_label(chrom)
    suffix = norm[3:] if norm.startswith("chr") else norm
    if suffix.isdigit():
        return (0, f"{int(suffix):02d}")
    if suffix == "X":
        return (1, "23")
    if suffix == "Y":
        return (1, "24")
    if suffix == "M":
        return (1, "25")
    return (2, suffix)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Force-genotype all cohort-observed targeted loci across every sample BAM, "
            "then export long-format and matrix summaries."
        )
    )
    parser.add_argument("--project-root", default=str(PATHS["project_root"]), help="Project root to scan for filtered VCFs")
    parser.add_argument("--reference", default=str(Path("/home/gadeaalonsoj/tfm/ref_alt/hg38_canonical.fa")), help="Reference FASTA used for mpileup")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Output directory for force-genotyping results")
    parser.add_argument("--threads", type=int, default=4, help="Threads to pass to bcftools output steps")
    parser.add_argument("--callable-dp", type=int, default=50, help="Depth threshold used to flag genotype calls as comfortably callable")
    parser.add_argument("--mpileup-max-depth", type=int, default=20000, help="High-depth cap for bcftools mpileup in this targeted panel")
    parser.add_argument("--keep-per-sample-vcfs", action="store_true", help="Retain per-sample force-genotyped VCFs instead of only summary tables")
    return parser.parse_args()


def discover_filtered_vcfs(project_root: Path) -> list[Path]:
    vcf_paths = sorted(project_root.glob("**/variants.filtered.vcf.gz"))
    if not vcf_paths:
        raise FileNotFoundError(f"No variants.filtered.vcf.gz files were found under {project_root}")
    return vcf_paths


def parse_bam_from_tvc_log(log_path: Path) -> Path:
    if not log_path.exists():
        raise FileNotFoundError(f"Missing TVC command log required to recover the sample BAM: {log_path}")
    line = log_path.read_text(encoding="utf-8", errors="replace").splitlines()[0].strip()
    if ":" in line:
        _, payload = line.split(":", 1)
    else:
        payload = line
    tokens = shlex.split(payload.strip())
    for idx, token in enumerate(tokens):
        if token == "--input-bam" and idx + 1 < len(tokens):
            bam_path = Path(tokens[idx + 1])
            if bam_path.exists():
                return bam_path
            raise FileNotFoundError(f"Recovered BAM path from {log_path} but the BAM is missing: {bam_path}")
    raise RuntimeError(f"Could not recover --input-bam from {log_path}")


def parse_sample_context(vcf_path: Path) -> dict[str, object]:
    sample_dir = vcf_path.parent
    parts = sample_dir.parts
    sample = sample_dir.name
    cohort = "Unknown"
    tissue = "Unknown"
    if len(parts) >= 3 and parts[-2] == "dna_calls":
        tissue = parts[-3]
        cohort = parts[-4] if len(parts) >= 4 else "Unknown"
    cohort = standardize_cohort_labels(pd.Series([cohort])).iloc[0].capitalize()
    tissue = standardize_tissue_labels(pd.Series([tissue])).iloc[0]
    bam_path = parse_bam_from_tvc_log(sample_dir / "logs" / "tvc.cmd.log")
    return {
        "sample": sample,
        "cohort": cohort,
        "tissue": tissue,
        "vcf_path": vcf_path,
        "sample_dir": sample_dir,
        "bam_path": bam_path,
    }


def query_vcf_sites(vcf_path: Path) -> list[dict[str, object]]:
    command = ["bcftools", "query", "-f", "%CHROM\t%POS\t%REF\t%ALT\n", str(vcf_path)]
    result = run_command(command, f"Querying loci from {vcf_path}")
    rows: list[dict[str, object]] = []
    for line in result.stdout.splitlines():
        if not line.strip():
            continue
        chrom, pos, ref, alt_field = line.split("\t")
        for alt in alt_field.split(","):
            alt = alt.strip()
            if not alt or alt == "." or alt == "<*>":
                continue
            rows.append(
                {
                    "CHROM": normalise_chromosome_label(chrom),
                    "POS": int(pos),
                    "REF": ref,
                    "ALT": alt,
                    "Variant_Key": f"{normalise_chromosome_label(chrom)}:{int(pos)}:{ref}:{alt}",
                    "Locus_Key": f"{normalise_chromosome_label(chrom)}:{int(pos)}:{ref}",
                }
            )
    return rows


def build_union_catalogues(sample_records: list[dict[str, object]]) -> tuple[pd.DataFrame, pd.DataFrame]:
    variant_rows: list[dict[str, object]] = []
    for record in sample_records:
        sample_sites = query_vcf_sites(Path(record["vcf_path"]))
        for site in sample_sites:
            variant_rows.append(
                {
                    **site,
                    "Sample": record["sample"],
                    "Cohort": record["cohort"],
                    "Tissue": record["tissue"],
                }
            )
    if not variant_rows:
        raise RuntimeError("No cohort-observed loci were recovered from the filtered VCF set.")

    variant_df = pd.DataFrame(variant_rows).drop_duplicates()
    variant_catalogue = (
        variant_df.groupby(["CHROM", "POS", "REF", "ALT", "Variant_Key", "Locus_Key"], as_index=False)
        .agg(
            Samples_with_variant=("Sample", "nunique"),
            Example_samples=("Sample", lambda s: "; ".join(sorted(s.astype(str).unique())[:6])),
            Cohorts_with_variant=("Cohort", lambda s: "; ".join(sorted(s.astype(str).unique()))),
            Tissues_with_variant=("Tissue", lambda s: "; ".join(sorted(s.astype(str).unique()))),
        )
        .sort_values(["CHROM", "POS", "REF", "ALT"], key=lambda col: col.map(chrom_sort_key) if col.name == "CHROM" else col)
        .reset_index(drop=True)
    )

    locus_catalogue = (
        variant_df.groupby(["CHROM", "POS", "REF", "Locus_Key"], as_index=False)
        .agg(
            Union_ALT_List=("ALT", lambda s: ",".join(sorted(dict.fromkeys(s.astype(str))))),
            N_Alt_Alleles=("ALT", lambda s: int(pd.Series(s.astype(str)).nunique())),
            Samples_with_any_variant=("Sample", "nunique"),
            Example_variant_samples=("Sample", lambda s: "; ".join(sorted(s.astype(str).unique())[:6])),
        )
        .sort_values(["CHROM", "POS", "REF"], key=lambda col: col.map(chrom_sort_key) if col.name == "CHROM" else col)
        .reset_index(drop=True)
    )
    locus_catalogue["Multi_Allelic_Union_Site"] = locus_catalogue["N_Alt_Alleles"] > 1
    return locus_catalogue, variant_catalogue


def write_candidate_vcf(locus_catalogue: pd.DataFrame, reference_fasta: Path, output_vcf: Path) -> Path:
    fai_path = Path(f"{reference_fasta}.fai")
    if not fai_path.exists():
        raise FileNotFoundError(f"Reference FASTA index is required but missing: {fai_path}")

    contig_lengths: dict[str, str] = {}
    for line in fai_path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        contig, length, *_ = line.split("\t")
        contig_lengths[normalise_chromosome_label(contig)] = length

    with output_vcf.open("w", encoding="utf-8") as handle:
        handle.write("##fileformat=VCFv4.2\n")
        for chrom in sorted(locus_catalogue["CHROM"].astype(str).unique(), key=chrom_sort_key):
            length = contig_lengths.get(chrom)
            if length:
                handle.write(f"##contig=<ID={chrom},length={length}>\n")
            else:
                handle.write(f"##contig=<ID={chrom}>\n")
        handle.write("#CHROM\tPOS\tID\tREF\tALT\tQUAL\tFILTER\tINFO\n")
        for row in locus_catalogue.itertuples(index=False):
            handle.write(f"{row.CHROM}\t{row.POS}\t.\t{row.REF}\t{row.Union_ALT_List}\t.\t.\t.\n")

    run_command(["bgzip", "-f", str(output_vcf)], f"Compressing candidate site VCF {output_vcf}")
    gz_path = output_vcf.with_suffix(output_vcf.suffix + ".gz")
    run_command(["tabix", "-f", "-p", "vcf", str(gz_path)], f"Indexing candidate site VCF {gz_path}")
    return gz_path


def parse_ad_field(ad_value: str | float | int | None) -> tuple[int | None, int | None]:
    if ad_value is None:
        return None, None
    text = str(ad_value).strip()
    if not text or text == "." or text.lower() == "nan":
        return None, None
    parts: list[int] = []
    for piece in text.split(","):
        piece = piece.strip()
        if not piece or piece == ".":
            continue
        try:
            parts.append(int(piece))
        except ValueError:
            return None, None
    if not parts:
        return None, None
    ref_depth = parts[0]
    alt_depth = sum(parts[1:]) if len(parts) > 1 else 0
    return ref_depth, alt_depth


def classify_raw_genotype(gt: str, dp: int | None) -> str:
    text = str(gt).strip()
    if not text or text in {".", "./.", ".|."} or dp is None or dp <= 0:
        return "No_call"
    alleles = re.split(r"[/|]", text)
    if not alleles or any(allele == "." or allele == "" for allele in alleles):
        return "No_call"
    try:
        numeric = [int(allele) for allele in alleles]
    except ValueError:
        return "No_call"
    nonzero = [allele for allele in numeric if allele > 0]
    if not nonzero:
        return "WT"
    if len(nonzero) == 1 and len(set(nonzero)) == 1:
        return "Het_ALT"
    if len(set(numeric)) == 1 and numeric[0] > 0:
        return "Hom_ALT"
    return "Other_ALT_model"


def thresholded_status(raw_class: str, dp: int | None, callable_dp: int) -> str:
    if raw_class == "No_call":
        return "No_call"
    if dp is None or dp < callable_dp:
        return "Below_DP_threshold"
    return raw_class


def query_force_genotype_vcf(vcf_path: Path) -> pd.DataFrame:
    query_fmt = "%CHROM\t%POS\t%REF\t%ALT[\t%GT\t%DP\t%AD]\n"
    result = run_command(["bcftools", "query", "-f", query_fmt, str(vcf_path)], f"Querying forced genotypes from {vcf_path}")
    rows: list[dict[str, object]] = []
    for line in result.stdout.splitlines():
        if not line.strip():
            continue
        fields = line.split("\t")
        if len(fields) < 7:
            continue
        chrom, pos, ref, alt, gt, dp, ad = fields[:7]
        dp_int = int(dp) if dp not in {".", ""} else None
        ref_depth, alt_depth = parse_ad_field(ad)
        alt_fraction = None
        if dp_int and alt_depth is not None:
            alt_fraction = alt_depth / dp_int if dp_int > 0 else None
        rows.append(
            {
                "CHROM": normalise_chromosome_label(chrom),
                "POS": int(pos),
                "REF": ref,
                "Observed_ALT_Call": "" if alt == "." else alt,
                "GT": gt,
                "DP": dp_int,
                "AD": ad,
                "Ref_Depth": ref_depth,
                "Alt_Depth": alt_depth,
                "Alt_Fraction": alt_fraction,
                "Locus_Key": f"{normalise_chromosome_label(chrom)}:{int(pos)}:{ref}",
            }
        )
    return pd.DataFrame(rows)


def force_genotype_sample(
    record: dict[str, object],
    candidate_vcf_gz: Path,
    reference_fasta: Path,
    output_dir: Path,
    threads: int,
    mpileup_max_depth: int,
    callable_dp: int,
    keep_per_sample_vcfs: bool,
    locus_catalogue: pd.DataFrame,
) -> pd.DataFrame:
    sample = str(record["sample"])
    bam_path = Path(record["bam_path"])
    sample_vcf = output_dir / "per_sample_vcf" / f"{sample}.forced_union_sites.vcf.gz"
    sample_log = output_dir / "logs" / f"{sample}.force_genotype.log"
    sample_table_path = output_dir / "per_sample_tables" / f"{sample}.tsv"
    ensure_directory(sample_vcf.parent)
    ensure_directory(sample_log.parent)
    ensure_directory(sample_table_path.parent)
    if sample_table_path.exists():
        return pd.read_csv(sample_table_path, sep="	")

    mpileup_cmd = [
        "bcftools",
        "mpileup",
        "-Ou",
        "-f",
        str(reference_fasta),
        "-T",
        str(candidate_vcf_gz),
        "-a",
        "AD,DP",
        "-d",
        str(mpileup_max_depth),
        str(bam_path),
    ]
    if threads > 1:
        mpileup_cmd.extend(["--threads", str(threads)])

    call_cmd = ["bcftools", "call", "-m", "-Oz", "-o", str(sample_vcf)]
    if threads > 1:
        call_cmd.extend(["--threads", str(threads)])

    with sample_log.open("wb") as log_handle:
        log_handle.write(("MPileup command: " + " ".join(mpileup_cmd) + "\n").encode())
        log_handle.write(("Call command: " + " ".join(call_cmd) + "\n").encode())
        mpileup_proc = subprocess.Popen(mpileup_cmd, stdout=subprocess.PIPE, stderr=log_handle)
        try:
            assert mpileup_proc.stdout is not None
            call_proc = subprocess.run(call_cmd, stdin=mpileup_proc.stdout, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, check=False)
        finally:
            if mpileup_proc.stdout is not None:
                mpileup_proc.stdout.close()
        mpileup_return = mpileup_proc.wait()
        if call_proc.stderr:
            log_handle.write(call_proc.stderr)

    if mpileup_return != 0 or call_proc.returncode != 0:
        raise RuntimeError(
            f"Force genotyping failed for {sample}. See {sample_log} for details. "
            f"mpileup_return={mpileup_return}, call_return={call_proc.returncode}"
        )

    run_command(["tabix", "-f", "-p", "vcf", str(sample_vcf)], f"Indexing forced genotype VCF for {sample}")
    sample_calls = query_force_genotype_vcf(sample_vcf)

    merged = locus_catalogue.merge(sample_calls, on=["CHROM", "POS", "REF", "Locus_Key"], how="left")
    merged["Sample"] = sample
    merged["Cohort"] = record["cohort"]
    merged["Tissue"] = record["tissue"]
    merged["Observed_ALT_Call"] = merged["Observed_ALT_Call"].fillna("")
    merged["GT"] = merged["GT"].fillna("./.")
    merged["DP"] = pd.to_numeric(merged["DP"], errors="coerce")
    merged["Ref_Depth"] = pd.to_numeric(merged["Ref_Depth"], errors="coerce")
    merged["Alt_Depth"] = pd.to_numeric(merged["Alt_Depth"], errors="coerce")
    merged["Alt_Fraction"] = pd.to_numeric(merged["Alt_Fraction"], errors="coerce")
    merged["AD"] = merged["AD"].fillna("")
    merged["Raw_Genotype_Class"] = [classify_raw_genotype(gt, None if pd.isna(dp) else int(dp)) for gt, dp in zip(merged["GT"], merged["DP"], strict=False)]
    merged["Genotype_Status"] = [thresholded_status(raw, None if pd.isna(dp) else int(dp), callable_dp) for raw, dp in zip(merged["Raw_Genotype_Class"], merged["DP"], strict=False)]
    merged["Meets_Callable_DP"] = merged["DP"].fillna(0).astype(float) >= callable_dp
    merged["Called_Nonref"] = merged["Raw_Genotype_Class"].isin(["Het_ALT", "Hom_ALT", "Other_ALT_model"])

    if not keep_per_sample_vcfs:
        for extra_path in [sample_vcf, Path(f"{sample_vcf}.tbi")]:
            if extra_path.exists():
                extra_path.unlink()

    column_order = [
        "Sample",
        "Cohort",
        "Tissue",
        "CHROM",
        "POS",
        "REF",
        "Union_ALT_List",
        "Observed_ALT_Call",
        "Locus_Key",
        "Multi_Allelic_Union_Site",
        "GT",
        "Raw_Genotype_Class",
        "Genotype_Status",
        "Meets_Callable_DP",
        "Called_Nonref",
        "DP",
        "AD",
        "Ref_Depth",
        "Alt_Depth",
        "Alt_Fraction",
        "N_Alt_Alleles",
        "Samples_with_any_variant",
    ]
    out = merged[column_order].copy()
    out.to_csv(sample_table_path, sep="	", index=False)
    return out


def summarise_loci(long_df: pd.DataFrame, locus_catalogue: pd.DataFrame) -> pd.DataFrame:
    def _count_status(status: str):
        return lambda series: int((series.astype(str) == status).sum())

    summary = (
        long_df.groupby(["CHROM", "POS", "REF", "Union_ALT_List", "Locus_Key", "Multi_Allelic_Union_Site", "N_Alt_Alleles", "Samples_with_any_variant"], as_index=False)
        .agg(
            N_Samples=("Sample", "nunique"),
            N_WT=("Genotype_Status", _count_status("WT")),
            N_Het_ALT=("Genotype_Status", _count_status("Het_ALT")),
            N_Hom_ALT=("Genotype_Status", _count_status("Hom_ALT")),
            N_Below_DP_Threshold=("Genotype_Status", _count_status("Below_DP_threshold")),
            N_No_Call=("Genotype_Status", _count_status("No_call")),
            N_Other_ALT_Model=("Genotype_Status", _count_status("Other_ALT_model")),
            Mean_DP=("DP", "mean"),
            Median_DP=("DP", "median"),
            ALT_Carrier_Samples=("Sample", lambda s: "; ".join(sorted(long_df.loc[s.index, "Sample"][long_df.loc[s.index, "Called_Nonref"]].astype(str).unique())[:8])),
            Cohorts_with_ALT=("Cohort", lambda s: "; ".join(sorted(long_df.loc[s.index, "Cohort"][long_df.loc[s.index, "Called_Nonref"]].astype(str).unique()))),
            Tissues_with_ALT=("Tissue", lambda s: "; ".join(sorted(long_df.loc[s.index, "Tissue"][long_df.loc[s.index, "Called_Nonref"]].astype(str).unique()))),
        )
        .sort_values(["CHROM", "POS", "REF"], key=lambda col: col.map(chrom_sort_key) if col.name == "CHROM" else col)
        .reset_index(drop=True)
    )
    summary["ALT_Carrier_Samples"] = summary["ALT_Carrier_Samples"].fillna("")
    summary["Cohorts_with_ALT"] = summary["Cohorts_with_ALT"].fillna("")
    summary["Tissues_with_ALT"] = summary["Tissues_with_ALT"].fillna("")
    summary["Mean_DP"] = summary["Mean_DP"].round(2)
    summary["Median_DP"] = summary["Median_DP"].round(2)
    return summary


def summarise_samples(long_df: pd.DataFrame) -> pd.DataFrame:
    sample_summary = (
        long_df.groupby(["Sample", "Cohort", "Tissue"], as_index=False)
        .agg(
            N_Loci=("Locus_Key", "nunique"),
            N_WT=("Genotype_Status", lambda s: int((s.astype(str) == "WT").sum())),
            N_Het_ALT=("Genotype_Status", lambda s: int((s.astype(str) == "Het_ALT").sum())),
            N_Hom_ALT=("Genotype_Status", lambda s: int((s.astype(str) == "Hom_ALT").sum())),
            N_Below_DP_Threshold=("Genotype_Status", lambda s: int((s.astype(str) == "Below_DP_threshold").sum())),
            N_No_Call=("Genotype_Status", lambda s: int((s.astype(str) == "No_call").sum())),
            Mean_DP=("DP", "mean"),
            Median_DP=("DP", "median"),
        )
        .sort_values(["Cohort", "Tissue", "Sample"], kind="stable")
        .reset_index(drop=True)
    )
    sample_summary["Mean_DP"] = sample_summary["Mean_DP"].round(2)
    sample_summary["Median_DP"] = sample_summary["Median_DP"].round(2)
    return sample_summary


def build_genotype_matrix(long_df: pd.DataFrame, sample_order: list[str]) -> pd.DataFrame:
    matrix = long_df.pivot_table(
        index=["CHROM", "POS", "REF", "Union_ALT_List", "Multi_Allelic_Union_Site", "Locus_Key"],
        columns="Sample",
        values="Genotype_Status",
        aggfunc="first",
        observed=False,
    )
    matrix = matrix.reindex(columns=sample_order)
    return matrix.reset_index()


def build_summary_table(
    sample_records: list[dict[str, object]],
    locus_catalogue: pd.DataFrame,
    variant_catalogue: pd.DataFrame,
    long_df: pd.DataFrame,
    callable_dp: int,
    mpileup_max_depth: int,
) -> pd.DataFrame:
    counts = Counter(long_df["Genotype_Status"].astype(str))
    rows = [
        ("Total samples force-genotyped", len(sample_records)),
        ("Total cohort-observed loci force-genotyped", int(locus_catalogue["Locus_Key"].nunique())),
        ("Total unique variant alleles represented in the union catalogue", int(variant_catalogue["Variant_Key"].nunique())),
        ("Multi-allelic union loci", int(locus_catalogue["Multi_Allelic_Union_Site"].sum())),
        ("WT genotype cells at callable depth", counts.get("WT", 0)),
        ("Het ALT genotype cells at callable depth", counts.get("Het_ALT", 0)),
        ("Hom ALT genotype cells at callable depth", counts.get("Hom_ALT", 0)),
        ("Genotype cells below the callable depth threshold", counts.get("Below_DP_threshold", 0)),
        ("No-call genotype cells", counts.get("No_call", 0)),
        ("Callable depth threshold used for interpretation", callable_dp),
        ("bcftools mpileup max-depth used for this deep targeted panel", mpileup_max_depth),
        ("Scope note", "Force-genotyped all loci observed at least once in the filtered targeted VCF set; this is not a gVCF of every base in the BED."),
        ("Interpretation note", "Main script-05 filtered calls remain the high-confidence discovery set. This workbook adds genotype context for WT / Het / Hom / no-call across cohort-observed loci."),
    ]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def style_workbook(workbook_path: Path) -> None:
    wb = load_workbook(workbook_path)
    for ws in wb.worksheets:
        max_col = ws.max_column
        max_row = ws.max_row
        if max_row >= 1:
            for cell in ws[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        for col_idx in range(1, max_col + 1):
            col_letter = get_column_letter(col_idx)
            width = 0
            for cell in ws[col_letter]:
                value = "" if cell.value is None else str(cell.value)
                width = max(width, min(len(value) + 2, 55))
                if cell.row > 1:
                    cell.alignment = WRAP_ALIGNMENT
            ws.column_dimensions[col_letter].width = max(12, width)

    if "README" in wb.sheetnames:
        ws = wb["README"]
        ws["A1"].font = TITLE_FONT

    for sheet_name in ["Per_Sample_Genotypes", "Genotype_Matrix"]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        status_column = None
        if sheet_name == "Per_Sample_Genotypes":
            for cell in ws[1]:
                if cell.value == "Genotype_Status":
                    status_column = cell.column
                    break
            if status_column:
                for row_idx in range(2, ws.max_row + 1):
                    value = str(ws.cell(row=row_idx, column=status_column).value or "")
                    fill = STATUS_FILLS.get(value)
                    if fill:
                        ws.cell(row=row_idx, column=status_column).fill = fill
        else:
            meta_headers = {"CHROM", "POS", "REF", "Union_ALT_List", "Multi_Allelic_Union_Site", "Locus_Key"}
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    header = ws.cell(row=1, column=cell.column).value
                    if header in meta_headers:
                        continue
                    fill = STATUS_FILLS.get(str(cell.value or ""))
                    if fill:
                        cell.fill = fill

    wb.save(workbook_path)


def main() -> None:
    args = parse_args()
    ensure_tool("bcftools")
    ensure_tool("bgzip")
    ensure_tool("tabix")

    project_root = Path(args.project_root).resolve()
    reference_fasta = Path(args.reference).resolve()
    output_dir = ensure_directory(Path(args.output_dir).resolve())
    if not reference_fasta.exists():
        raise FileNotFoundError(f"Reference FASTA not found: {reference_fasta}")

    log(f"Scanning filtered targeted VCFs under {project_root}")
    vcf_paths = discover_filtered_vcfs(project_root)
    sample_records = [parse_sample_context(vcf_path) for vcf_path in vcf_paths]
    sample_records = sorted(sample_records, key=lambda record: (str(record["cohort"]), str(record["tissue"]), str(record["sample"])))
    log(f"Recovered {len(sample_records)} filtered VCFs / sample BAM pairs")

    locus_catalogue, variant_catalogue = build_union_catalogues(sample_records)
    log(
        "Built the cohort-wide target catalogue: "
        f"{len(locus_catalogue)} loci and {len(variant_catalogue)} unique variant alleles"
    )

    candidate_vcf_gz = write_candidate_vcf(locus_catalogue, reference_fasta, output_dir / "cohort_union_sites.vcf")
    log(f"Wrote the cohort union site VCF to {candidate_vcf_gz}")

    per_sample_tables: list[pd.DataFrame] = []
    for index, record in enumerate(sample_records, start=1):
        log(f"Force genotyping sample {index}/{len(sample_records)}: {record['sample']}")
        per_sample_tables.append(
            force_genotype_sample(
                record=record,
                candidate_vcf_gz=candidate_vcf_gz,
                reference_fasta=reference_fasta,
                output_dir=output_dir,
                threads=args.threads,
                mpileup_max_depth=args.mpileup_max_depth,
                callable_dp=args.callable_dp,
                keep_per_sample_vcfs=args.keep_per_sample_vcfs,
                locus_catalogue=locus_catalogue,
            )
        )

    long_df = pd.concat(per_sample_tables, ignore_index=True)
    long_df = long_df.sort_values(["Cohort", "Tissue", "Sample", "CHROM", "POS"], key=lambda col: col.map(chrom_sort_key) if col.name == "CHROM" else col, kind="stable").reset_index(drop=True)
    locus_summary = summarise_loci(long_df, locus_catalogue)
    sample_summary = summarise_samples(long_df)
    sample_order = sample_summary["Sample"].astype(str).tolist()
    genotype_matrix = build_genotype_matrix(long_df, sample_order)
    summary_table = build_summary_table(sample_records, locus_catalogue, variant_catalogue, long_df, args.callable_dp, args.mpileup_max_depth)

    readme_table = pd.DataFrame(
        {
            "Item": [
                "What this stage does",
                "How it differs from script 05",
                "What union sites means here",
                "How to read Genotype_Status",
                "Key limitation",
            ],
            "Value": [
                "Force-genotypes every cohort-observed targeted locus across all sample BAMs so WT / Het / Hom / no-call states can be examined directly.",
                "Script 05 remains the strict discovery call set. This stage adds genotype context and does not replace the filtered PASS-based VCFs.",
                "A union site is any targeted locus observed at least once anywhere in the filtered cohort VCF set.",
                f"WT / Het_ALT / Hom_ALT require DP >= {args.callable_dp}; lower-depth loci are labelled Below_DP_threshold and missing-depth loci are labelled No_call.",
                "This is not a full gVCF of every base in the BED. Loci absent from the union catalogue were never observed in the original filtered cohort call set.",
            ],
        }
    )

    workbook_path = output_dir / "GSDMB_Forced_Genotypes_Union_Sites.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        readme_table.to_excel(writer, sheet_name="README", index=False)
        summary_table.to_excel(writer, sheet_name="Summary", index=False)
        locus_summary.to_excel(writer, sheet_name="Locus_Summary", index=False)
        sample_summary.to_excel(writer, sheet_name="Sample_Summary", index=False)
        long_df.to_excel(writer, sheet_name="Per_Sample_Genotypes", index=False)
        genotype_matrix.to_excel(writer, sheet_name="Genotype_Matrix", index=False)
        variant_catalogue.to_excel(writer, sheet_name="Variant_Catalogue", index=False)

    style_workbook(workbook_path)

    locus_catalogue.to_csv(output_dir / "05b_union_locus_catalogue.tsv", sep="\t", index=False)
    variant_catalogue.to_csv(output_dir / "05b_union_variant_catalogue.tsv", sep="\t", index=False)
    long_df.to_csv(output_dir / "05b_per_sample_force_genotypes.tsv", sep="\t", index=False)
    locus_summary.to_csv(output_dir / "05b_locus_summary.tsv", sep="\t", index=False)
    sample_summary.to_csv(output_dir / "05b_sample_summary.tsv", sep="\t", index=False)
    genotype_matrix.to_csv(output_dir / "05b_genotype_matrix.tsv", sep="\t", index=False)
    summary_table.to_csv(output_dir / "05b_summary.tsv", sep="\t", index=False)

    log(f"Force genotyping completed successfully. Workbook: {workbook_path}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:  # pragma: no cover - explicit user-facing failure path
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
